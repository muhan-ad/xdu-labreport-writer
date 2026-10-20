"""Report-specific prose, live-bound Hall values and safe fallback from old guidance."""
import contextlib
import importlib.util
import io
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'物理实验/实验脚本'
sys.path.insert(0,str(BASE))
from common.variants import compose

class ReportProseTests(unittest.TestCase):
    def test_old_polish_guidance_falls_back_to_experiment_prose(self):
        values={'b2_min':1,'b2_max':2,'b2_span':1,'k':1,'b2_mean':1}
        override={'误差分析':'误差应结合本次数据和仪器说明评定。重复测量的随机分散用于A类不确定度。'}
        with patch.dict(os.environ,{'LAB_POLISH':json.dumps(override,ensure_ascii=False)},clear=False):
            with patch.dict(os.environ,{},clear=False),contextlib.redirect_stdout(io.StringIO()):
                old=os.environ.pop('LAB_JOB_INPUT',None)
                try:result=compose(str(BASE/'霍尔效应实验'),values)
                finally:
                    if old is not None:os.environ['LAB_JOB_INPUT']=old
        self.assertNotIn('误差应结合',result['误差分析'])
        self.assertIn('1～2 T',result['误差分析'])
        self.assertNotIn('%%DATA:',result['误差分析'])

    def test_all_builtin_variants_have_no_software_instructions(self):
        from common.content_policy import contains_writing_instructions
        def visit(value):
            if isinstance(value, str):
                self.assertFalse(contains_writing_instructions(value), value)
                self.assertNotIn('误差在合理范围', value)
            elif isinstance(value, list):
                for item in value: visit(item)
            elif isinstance(value, dict):
                for item in value.values(): visit(item)
        for path in BASE.glob('*/variants.json'):
            with self.subTest(experiment=path.parent.name):
                visit(json.loads(path.read_text(encoding='utf-8')))

    def test_all_generated_reports_have_no_software_instructions(self):
        from common.content_policy import contains_writing_instructions
        sys.path.insert(0, str(ROOT/'tests'))
        from custom_plot_test import all_experiments, run_report
        from docx import Document
        forbidden=('结果在合理误差范围内', '实验数据可靠', '线性关系良好', '验证了', '实验流程和数据处理基本合理')
        for exp in all_experiments():
            with self.subTest(experiment=exp), contextlib.redirect_stdout(io.StringIO()):
                path=run_report(exp)
                try:
                    doc=Document(path)
                    text='\n'.join(n.text or '' for n in doc.element.iter() if n.tag.endswith('}t'))
                    self.assertFalse(contains_writing_instructions(text), exp)
                    self.assertNotIn('%%DATA:', text)
                    for phrase in forbidden: self.assertNotIn(phrase, text)
                finally: shutil.rmtree(Path(path).parent)

    def test_real_hall_generation_uses_current_values_without_guidance(self):
        exp=BASE/'霍尔效应实验'
        spec=importlib.util.spec_from_file_location('hall_prose_generate',exp/'generate.py')
        module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        data=json.loads((ROOT/'tests/fixtures/experiment-data.json').read_text(encoding='utf-8'))['霍尔效应实验']
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary);shutil.copy2(exp/'variants.json',root/'variants.json');shutil.copytree(exp/'rag',root/'rag')
            module.SCRIPT_DIR=str(root)
            from docx import Document
            with patch.dict(os.environ,{},clear=True),contextlib.redirect_stdout(io.StringIO()):
                module._generate_docx(data,str(root/'hall.docx'))
            doc=Document(root/'hall.docx');text='\n'.join(p.text for p in doc.paragraphs)
            self.assertNotIn('误差应结合',text);self.assertNotIn('%%DATA:',text)
            expected=[module._hall_voltage(u1,u2)/(data['kh']*i) for u1,u2,i in zip(data['t2_u1'],data['t2_u2'],data['t2_i'])]
            self.assertIn(f'{min(expected):.6g}～{max(expected):.6g} T',text)
            self.assertEqual(len(doc.inline_shapes),2)

if __name__=='__main__':unittest.main()
