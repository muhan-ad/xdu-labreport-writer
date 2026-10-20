"""官网预设：真实测试模板只读；数值、配对、组图及原子插入回归。"""
import copy
import importlib.util
import json
import math
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'物理实验/实验脚本'
sys.path.insert(0,str(BASE))
from common.official_charts import chart_specs,render_specs

def data(name):return json.loads((BASE/name/'test-data.json').read_text(encoding='utf-8-sig'))
def specs(name,values=None,key='all'):return chart_specs(name,values if values is not None else data(name),str(BASE),key)

class OfficialChartTests(unittest.TestCase):
 def test_all_dedicated_groups_render_without_mutating_data(self):
  names=['RLC电路的稳态特性研究实验','光的偏振特性测量','光的偏振特性测量（第二学期）','电容与高电阻的测量','衍射光强分布的测量','灵敏电流计特性的测量','电表的改装与校准','电子偏转特性的测量','理想气体状态方程','霍尔效应实验','重力加速度的测量','直螺线管磁场分布的测量','静电场的模拟']
  with tempfile.TemporaryDirectory() as temp:
   for i,name in enumerate(names):
    with self.subTest(name=name):
     d=data(name);before=copy.deepcopy(d);charts=specs(name,d)
     images=render_specs(charts,str(Path(temp)/f'{i}.png'),str(BASE))
     self.assertEqual(d,before);self.assertEqual(len(images),len(charts))
     self.assertTrue(all(Path(img['path']).stat().st_size>1000 for img in images))
     self.assertTrue(all(img['caption'] and img['note'] for img in images))
 def test_axes_values_and_correction_sign(self):
  name='电表的改装与校准';d=data(name);s=specs(name,d)
  self.assertEqual(s[0]['series'][0]['x'],d['i_x'])
  self.assertAlmostEqual(s[0]['series'][0]['y'][0],d['i_s'][0]-d['i_x'][0])
  d=data('光的偏振特性测量');s=specs('光的偏振特性测量',d)[0]
  self.assertEqual(s['series'][0]['y'],d['p_values'])
  s=specs('衍射光强分布的测量')[0];self.assertEqual(max(s['series'][0]['y']),1)
  s=specs('电子偏转特性的测量');self.assertEqual(len(s),3)
  self.assertEqual([a['label'] for a in s[0]['series']],['加速电压1000 V','加速电压800 V'])
  d=data('理想气体状态方程');s=specs('理想气体状态方程',d)
  self.assertAlmostEqual(s[0]['series'][0]['x'][0],1/(d['p0']+d['ta_dp'][0]))
  self.assertAlmostEqual(s[1]['series'][0]['x'][0],d['tb_t'][0]+273.15)
  d=data('霍尔效应实验');s=specs('霍尔效应实验',d)
  self.assertEqual(len(s),2);self.assertEqual(s[0]['ylabel'],'UH / mV')
  b=specs('霍尔效应实验',d,'field_i')[0]
  self.assertAlmostEqual(b['series'][0]['y'][0],(abs(d['t2_u1'][0])+abs(d['t2_u2'][0]))/2/(d['kh']*d['t2_i'][0]))
 def test_pairing_never_truncates_or_filters_independently(self):
  name='电表的改装与校准';d=data(name);d['i_s']=d['i_s'][:-1]
  with self.assertRaisesRegex(ValueError,'长度不匹配'):specs(name,d)
  d=data(name);d['i_s'][2]=None
  with self.assertRaisesRegex(ValueError,'第3点'):specs(name,d)
  d=data(name);d['i_x'][2]=None;d['i_s'][2]=None
  self.assertEqual(len(specs(name,d,'ammeter')[0]['series'][0]['x']),len(d['i_x'])-1)
  d={'t_values':[0,1,2],'q_values':[1,0,.5]}
  with self.assertRaisesRegex(ValueError,'大于0'):specs('电容与高电阻的测量',d)
  with self.assertRaisesRegex(ValueError,'尚无'):specs('三棱镜顶角的测量',{})
 def test_selected_chart_does_not_require_other_group(self):
  s=specs('电表的改装与校准',{'i_x':[.1,.2,.3],'i_s':[.11,.21,.29]},'ammeter')
  self.assertEqual(len(s),1)
 def test_rlc_missing_half_power_coverage_is_not_invented(self):
  d=data('RLC电路的稳态特性研究实验')
  for row in d['t1']:row[2]=5
  s=specs('RLC电路的稳态特性研究实验',d,'amplitude')[0]
  self.assertEqual(s['bandwidth'],[None,None]);self.assertIn('未完整覆盖',s['note'])
 def test_cli_multi_image_and_atomic_report_group(self):
  from docx import Document
  from docx.enum.text import WD_ALIGN_PARAGRAPH
  helper=ROOT/'src/main/chart-preview.py';insert=ROOT/'src/main/insert-chart-into-docx.py'
  with tempfile.TemporaryDirectory() as temp:
   exp=Path(temp)/'experiment';exp.mkdir();name='RLC电路的稳态特性研究实验'
   for file in ['schema.json','test-data.json']:
    (exp/('data.json' if file=='test-data.json' else file)).write_bytes((BASE/name/file).read_bytes())
   output=Path(temp)/'plot.png'
   child_env={**os.environ,'PYTHONUTF8':'1','PYTHONIOENCODING':'utf-8'}
   run=subprocess.run([sys.executable,str(helper),'--exp-path',str(exp),'--x-field','__official__','--y-field','__official__','--chart-type','official','--scripts-root',str(BASE),'--output',str(output)],capture_output=True,encoding='utf-8',env=child_env)
   self.assertEqual(run.returncode,0,run.stderr);result=json.loads(run.stdout.strip().splitlines()[-1]);self.assertEqual(len(result['images']),2)
   report=Path(temp)/'report.docx';doc=Document();doc.add_heading('实验结果分析',1);doc.add_paragraph('原始正文');doc.add_heading('误差分析',1);doc.save(report)
   original=report.read_bytes()
   bad=copy.deepcopy(result['images']);bad[1]['path']=str(Path(temp)/'missing.png')
   args=[sys.executable,str(insert),'--docx-path',str(report),'--image-path',str(output),'--images-json']
   fail=subprocess.run(args+[json.dumps(bad,ensure_ascii=False)],capture_output=True,env=child_env)
   self.assertNotEqual(fail.returncode,0);self.assertEqual(report.read_bytes(),original)
   ok=subprocess.run(args+[json.dumps(result['images'],ensure_ascii=False)],capture_output=True,env=child_env)
   self.assertEqual(ok.returncode,0,ok.stderr.decode('utf-8'))
   doc=Document(report);self.assertEqual(len(doc.inline_shapes),2)
   captions=[p for p in doc.paragraphs if p.text.startswith('图')]
   self.assertEqual(len(captions),2);self.assertTrue(all(p.alignment==WD_ALIGN_PARAGRAPH.CENTER for p in captions))
   self.assertTrue(any(p.text.startswith('注：') for p in doc.paragraphs))

if __name__=='__main__':unittest.main()
