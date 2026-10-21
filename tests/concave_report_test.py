"""Second-semester lens record, independent arithmetic and old-data compatibility."""
import contextlib
import copy
import importlib.util
import io
import json
import math
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
EXP=ROOT/'物理实验/实验脚本/薄透镜焦距的测量（凹透镜）'
sys.path.insert(0,str(EXP.parent))
spec=importlib.util.spec_from_file_location('concave_report',EXP/'generate.py')
gen=importlib.util.module_from_spec(spec);spec.loader.exec_module(gen)

class ConcaveReportTests(unittest.TestCase):
    def setUp(self):
        self.data=json.loads((EXP/'test-data.json').read_text(encoding='utf-8'))
    def test_independent_convex_mean_and_uncertainty(self):
        r=gen._compute(self.data)
        # Eight symmetric readings: mean=15 cm, sample variance=0.0004 cm².
        self.assertAlmostEqual(r['f_auto_bar'],15.0,places=12)
        self.assertAlmostEqual(r['convex']['s'],0.02,places=12)
        self.assertAlmostEqual(r['convex']['uA'],0.02/math.sqrt(8),places=12)
        self.assertAlmostEqual(r['convex']['uB'],0.05/math.sqrt(3),places=12)
        self.assertAlmostEqual(r['u_auto'],math.sqrt(0.0004/8+0.05**2/3),places=12)
    def test_independent_position_difference_and_thin_lens_oracles(self):
        r=gen._compute(self.data)
        expected=[b-a for a,b in zip(self.data['x6'],self.data['x7'])]
        self.assertEqual(r['ca_gap'],expected)
        self.assertAlmostEqual(r['f_ca_abs_bar'],sum(expected)/8,places=12)
        self.assertAlmostEqual(r['uB_ca'],math.sqrt(2*0.05**2/3),places=12)
        # First record: u=6.02, v=10.03, f=60.3806/-4.01 cm.
        self.assertAlmostEqual(r['f_im_vals'][0],-60.3806/4.01,places=12)
        self.assertTrue(all(f<0 for f in r['f_im_vals']))
        self.assertEqual(self.data['concave'][0],[24.02,18.0,28.03])
    def test_old_records_preserved_and_new_measurement_never_fabricated(self):
        legacy=copy.deepcopy(self.data);legacy.pop('f_auto');before=copy.deepcopy(legacy)
        with self.assertRaisesRegex(ValueError,'凸透镜焦距'):
            gen._compute(legacy)
        self.assertEqual(legacy,before)
    def test_incomplete_wrong_length_and_nonfinite_new_table_rejected(self):
        for values in [None,[15.0]*7,[15.0]*9,[15.0]*7+[None],[15.0]*7+[float('nan')],[15.0]*7+[False],[0.0]*8]:
            with self.subTest(values=values):
                self.data['f_auto']=values
                with self.assertRaises(ValueError):gen._compute(self.data)
    def test_impossible_optical_geometry_rejected_before_division(self):
        for row in [[24,18,24],[17,18,25],[30,18,25]]:
            with self.subTest(row=row):
                self.data['concave'][0]=row
                with self.assertRaisesRegex(ValueError,'表3'):gen._compute(self.data)
    def test_report_three_tables_numbering_results_and_no_false_agreement(self):
        with tempfile.TemporaryDirectory() as d,contextlib.redirect_stdout(io.StringIO()):
            out=Path(d)/'report.docx'
            with patch.dict('os.environ',{'LAB_POLISH':'','LAB_CUSTOM_QUIZ':'','LAB_JOB_INPUT':'','LAB_DATA_PHOTO':'','LAB_VARIANT_SELECTION':''}):
                gen._generate_docx(self.data,str(out))
            doc=Document(out)
            self.assertEqual([(len(t.rows),len(t.columns)) for t in doc.tables],[(2,9),(4,9),(4,7)])
            xml=''.join(doc.element.itertext())
            for phrase in ['表1 自准法测量凸透镜焦距','表2 自准法测量凹透镜焦距','表3 物像法测量凹透镜焦距','15.01','独立','只计算三次焦距的平均值']:
                self.assertIn(phrase,xml)
            for phrase in ['处于本实验的不确定度水平之内','说明测量自洽','%%DATA:']:
                self.assertNotIn(phrase,xml)
            cells=[''.join(c._tc.itertext()) for row in doc.tables[1].rows for c in row.cells]
            self.assertIn('9.02',''.join(cells));self.assertIn('24.03',''.join(cells))
            self.assertTrue(doc.element.xpath('.//m:oMath'))

if __name__=='__main__':unittest.main()
