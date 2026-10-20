"""Final OOXML layout covers incremental writes, lists, equations and photo precedence."""
import contextlib,io,json,os,re,shutil,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
ROOT=Path(__file__).resolve().parents[1];BASE=ROOT/'物理实验/实验脚本'
sys.path.insert(0,str(BASE));sys.path.insert(0,str(ROOT/'tests'))
from common.docx_report import DocxReportWriter
from custom_plot_test import all_experiments,run_report

def equations(p):
 return [n for n in p._p.findall('.//'+qn('m:oMath')) if re.search('[=≈≤≥∝∼≡]', ''.join(n.itertext()))]

class LayoutTests(unittest.TestCase):
 def test_every_builtin_variant_uses_final_layout_pass(self):
  def strings(value):
   if isinstance(value,str):yield value
   elif isinstance(value,dict):
    for v in value.values():yield from strings(v)
   elif isinstance(value,list):
    for v in value:yield from strings(v)
  with tempfile.TemporaryDirectory() as tmp,contextlib.redirect_stdout(io.StringIO()):
   for p in BASE.glob('*/variants.json'):
    with self.subTest(experiment=p.parent.name):
     out=Path(tmp)/'variant.docx';w=DocxReportWriter(str(out))
     for value in strings(json.loads(p.read_text(encoding='utf-8'))):w.add_paragraph_rich(re.sub(r'%%DATA:[^:]+:.+?%%','1.234',value))
     w.close();doc=Document(out)
     for paragraph in doc.paragraphs:self.assertLessEqual(len(equations(paragraph)),1,paragraph.text)
 def test_incremental_points_and_parallel_equations_preserve_native_math(self):
  with tempfile.TemporaryDirectory() as tmp:
   out=Path(tmp)/'report.docx';w=DocxReportWriter(str(out))
   w.add_paragraph('（1）读数误差');w.add_run('（2）温度误差');w.add_run('（10）接触误差')
   w.add_paragraph_rich('1. 甲项 2. 乙项')
   w.add_paragraph_rich('（一）第一原因（二）第二原因')
   w.add_paragraph_rich('1.原因甲；2.原因乙')
   w.add_paragraph_rich('理想调制度为1、1/2、0；式（1）代入式（2）。')
   w.add_paragraph('结果：');w.add_inline_math('a=1');w.add_run('；');w.add_inline_math('b=2')
   w.add_math(r'c=3;d=4');w.add_math(r'e=5,\quad f=6')
   w.add_paragraph_rich('公式 $g=7$；$h=8$。')
   w.add_math(r'$j=1;D_e\propto U;D_m\propto I$')
   w.save();w.close();doc=Document(out)
   for label in ['（1）读数误差','（2）温度误差','（10）接触误差','1. 甲项','2. 乙项']:
    self.assertIn(label,[p.text.strip() for p in doc.paragraphs])
   self.assertIn('理想调制度为1、1/2、0；式（1）代入式（2）。',[p.text for p in doc.paragraphs])
   self.assertIn('（二）第二原因',[p.text for p in doc.paragraphs])
   self.assertEqual(sum(len(equations(p)) for p in doc.paragraphs),11)
   self.assertTrue(all(len(equations(p))<=1 for p in doc.paragraphs))
   for p in doc.paragraphs:
    if equations(p):self.assertEqual(p.alignment,WD_ALIGN_PARAGRAPH.CENTER)
   # Native equation contents survive; no text substitutions or lost equalities.
   text=''.join(doc._element.itertext())
   self.assertNotIn('$',text)
   for symbol in 'abcdefgh':self.assertIn(symbol,text)
 def test_photo_suppresses_only_explicit_raw_tables(self):
  with tempfile.TemporaryDirectory() as tmp:
   photo=Path(tmp)/'record.png';Image.new('RGB',(100,80),'white').save(photo)
   for has_photo in [False,True]:
    out=Path(tmp)/f'{has_photo}.docx'
    with patch.dict(os.environ,{'LAB_DATA_PHOTO':str(photo) if has_photo else str(photo)+'.missing'}):
     w=DocxReportWriter(str(out));w.add_data_photo()
     w.add_table(['读数'],[['1']],role='raw');w.add_table(['计算结果'],[['2']]);w.close()
    self.assertEqual(len(Document(out).tables),1 if has_photo else 2)
 def test_all_reports_have_one_equation_per_paragraph_and_no_raw_only_duplicates(self):
  with tempfile.TemporaryDirectory() as tmp:
   photo=Path(tmp)/'record.png';Image.new('RGB',(100,80),'white').save(photo)
   for exp in all_experiments():
    with self.subTest(experiment=exp),patch.dict(os.environ,{'LAB_DATA_PHOTO':str(photo)}),contextlib.redirect_stdout(io.StringIO()):
     out=Path(run_report(exp));doc=Document(out)
     try:
      for p in doc.paragraphs:self.assertLessEqual(len(equations(p)),1,p.text)
      if exp=='静电场的模拟':self.assertEqual(len(doc.tables),1)
      if exp=='灵敏电流计特性的测量':self.assertEqual(len(doc.tables),1)
      if exp=='长度与体积的测量':self.assertEqual(len(doc.tables),2)
      if exp=='光的偏振特性测量（第二学期）':self.assertEqual(len(doc.tables),1)
      if exp=='重力加速度的测量':self.assertEqual(len(doc.tables[0].rows),4)
     finally:shutil.rmtree(out.parent)

if __name__=='__main__':unittest.main()
