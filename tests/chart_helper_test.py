# -*- coding: utf-8 -*-
"""图表辅助脚本协议测试（src/main/chart-preview.py、src/main/insert-chart-into-docx.py）。

两个脚本都由主进程以 **stdin 注入源码** 执行（打包后 app.asar 里的文件不是真实路径，
不能作为 python 入口），所以这里也按同样方式跑：

    python -B -X utf8 - <参数…>        （源码从 stdin 喂进去）

覆盖：
- 预览：真实实验的 data.json → 出 PNG + stdout JSON；字段不存在 → 中文报错 + 非 0 退出
- 插图：造一份带章节标题的 docx → 图片插进指定章节（图片数 +1、位置在目标章节内）；
  临时文件被清理；章节不存在 → 中文报错
"""

import json
import os
import subprocess
import sys
import tempfile

from docx import Document
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SCRIPTS_ROOT = os.path.join(ROOT, '物理实验', '实验脚本')
PREVIEW_PY = os.path.join(ROOT, 'src', 'main', 'chart-preview.py')
INSERT_PY = os.path.join(ROOT, 'src', 'main', 'insert-chart-into-docx.py')
EXP_NAME = '长度与体积的测量'          # 真实实验，data.json 里有等长数组字段
EXP_PATH = os.path.join(SCRIPTS_ROOT, EXP_NAME)


def run_helper(script, args, timeout=120):
    """按应用的方式跑辅助脚本：源码经 stdin 注入，参数走 argv。"""
    with open(script, encoding='utf-8') as f:
        source = f.read()
    proc = subprocess.run(
        [sys.executable, '-B', '-X', 'utf8', '-', *args],
        input=source, capture_output=True, text=True, encoding='utf-8', timeout=timeout)
    payload = None
    for text in (proc.stdout or '', proc.stderr or ''):
        lines = [l for l in text.strip().split('\n') if l.strip()]
        if lines:
            try:
                payload = json.loads(lines[-1])
                break
            except Exception:
                continue
    return proc, payload


def make_docx(path, headings=('一、实验数据记录', '二、数据处理', '三、实验结果分析', '四、思考题')):
    """造一份最小报告：章节标题 + 每节一句话，便于断言插图位置。"""
    doc = Document()
    for h in headings:
        doc.add_heading(h, level=1)
        doc.add_paragraph('这一节的内容。')
    doc.save(path)


def section_span(path, section):
    """返回 (标题段落下标, 下一章节标题下标)，用于判断图片落在哪一节里。"""
    paras = [p.text.strip() for p in Document(path).paragraphs]
    start = next((i for i, t in enumerate(paras) if t.endswith(section)), None)
    if start is None:
        return None
    end = next((i for i in range(start + 1, len(paras)) if paras[i].startswith(('一、', '二、', '三、', '四、', '五、'))), len(paras))
    return start, end


def image_paragraph_indexes(path):
    """含图片的段落下标（图片是 w:drawing）。"""
    out = []
    for i, p in enumerate(Document(path).paragraphs):
        if p._element.findall('.//' + qn('w:drawing')):
            out.append(i)
    return out


def main():
    problems = []
    work = tempfile.mkdtemp(prefix='chart_helper_')

    # ── 1) 预览：真实实验数据出图 ──
    out_png = os.path.join(work, 'preview.png')
    proc, payload = run_helper(PREVIEW_PY, [
        '--exp-path', EXP_PATH, '--x-field', 'D', '--y-field', 'd_shi',
        '--chart-type', 'scatter', '--output', out_png, '--scripts-root', SCRIPTS_ROOT,
    ])
    if not payload or not payload.get('ok'):
        problems.append('预览：应成功，实得 %s（stderr 尾部：%s）' % (payload, (proc.stderr or '')[-200:]))
    elif not os.path.exists(payload.get('path') or ''):
        problems.append('预览：回报的 PNG 不存在：%s' % payload.get('path'))
    elif os.path.getsize(payload['path']) < 1024:
        problems.append('预览：PNG 过小（%d 字节），可能没画出来' % os.path.getsize(payload['path']))
    elif not payload.get('title'):
        problems.append('预览：没有返回默认标题')
    else:
        print('   ✔ 预览：出图 %d 字节，标题「%s」' % (os.path.getsize(payload['path']), payload['title']))

    # 拟合/柱状也要能出图（chart_type 分支）
    for kind in ('fit', 'bar'):
        proc, payload = run_helper(PREVIEW_PY, [
            '--exp-path', EXP_PATH, '--x-field', 'D', '--y-field', 'd_shi',
            '--chart-type', kind, '--output', os.path.join(work, kind + '.png'),
            '--scripts-root', SCRIPTS_ROOT,
        ])
        if not payload or not payload.get('ok'):
            problems.append('预览(%s)：应成功，实得 %s（stderr 尾部：%s）' % (kind, payload, (proc.stderr or '')[-160:]))
    print('   ✔ 预览：scatter / fit / bar 三种类型都能出图')

    # 字段不存在 → 中文报错 + 非 0 退出
    proc, payload = run_helper(PREVIEW_PY, [
        '--exp-path', EXP_PATH, '--x-field', '不存在的字段', '--y-field', 'd_shi',
        '--output', os.path.join(work, 'bad.png'), '--scripts-root', SCRIPTS_ROOT,
    ])
    if proc.returncode == 0 or not payload or payload.get('ok') is not False:
        problems.append('预览：字段不存在应失败，实得 rc=%s payload=%s' % (proc.returncode, payload))
    elif '字段不存在' not in (payload.get('error') or ''):
        problems.append('预览：错误信息应说明字段不存在，实得 %s' % payload.get('error'))

    # ── 2) 插图：插进指定章节 ──
    docx_path = os.path.join(work, '报告.docx')
    make_docx(docx_path)
    before = len(image_paragraph_indexes(docx_path))
    proc, payload = run_helper(INSERT_PY, [
        '--docx-path', docx_path, '--image-path', out_png,
        '--section', '实验结果分析', '--width-cm', '12',
    ])
    if not payload or not payload.get('ok'):
        problems.append('插图：应成功，实得 %s（stderr 尾部：%s）' % (payload, (proc.stderr or '')[-200:]))
    else:
        after = image_paragraph_indexes(docx_path)
        if len(after) != before + 1:
            problems.append('插图：图片数应为 %d，实得 %d' % (before + 1, len(after)))
        else:
            span = section_span(docx_path, '实验结果分析')
            if not span or not (span[0] < after[0] < span[1]):
                problems.append('插图：图片不在「实验结果分析」章节内（图片段 %s，章节 %s）' % (after, span))
            elif not os.path.exists(docx_path + '.~chart.tmp'):
                print('   ✔ 插图：图片落在「实验结果分析」章节内，临时文件已清理')
            else:
                problems.append('插图：临时文件 .~chart.tmp 未清理')
        if payload.get('section') != '实验结果分析':
            problems.append('插图：回报章节名不对：%s' % payload.get('section'))

    # 章节不存在 → 中文报错
    docx2 = os.path.join(work, '报告2.docx')
    make_docx(docx2, headings=('一、实验数据记录', '二、数据处理'))
    proc, payload = run_helper(INSERT_PY, [
        '--docx-path', docx2, '--image-path', out_png, '--section', '实验结果分析',
    ])
    if proc.returncode == 0 or not payload or payload.get('ok') is not False:
        problems.append('插图：章节不存在应失败，实得 rc=%s payload=%s' % (proc.returncode, payload))
    elif '未找到章节' not in (payload.get('error') or ''):
        problems.append('插图：错误信息应说明未找到章节，实得 %s' % payload.get('error'))

    if problems:
        print('\n验证失败：')
        for p in problems:
            print('  - ' + p)
        return 1
    print('\nPASS: 图表辅助脚本（预览出图 / 插图定位 / 错误回报 / stdin 注入可运行）')
    return 0


if __name__ == '__main__':
    sys.exit(main())
