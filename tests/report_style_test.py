"""All experiments: numbered error paragraphs and centered titles below pictures."""
import contextlib, importlib.util, io, json, os, re, shutil, sys, tempfile, unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parents[1];BASE=ROOT/'物理实验/实验脚本'
sys.path.insert(0,str(BASE));sys.path.insert(0,str(ROOT/'tests'))
from common.report_style import error_points
from common.docx_report import DocxReportWriter
from custom_plot_test import all_experiments,run_report
class ReportStyleTests(unittest.TestCase):
    def test_error_numbering_protects_formulas_and_is_idempotent(self):
        for value in ['读数误差影响结果。关系为 $u=1.23$。结果 %%DATA:x:%.2f%%。', '（1）读数误差。\n（2）温度误差。', '1. 读数误差。\n2. 温度误差。']:
            rendered=error_points(value)
            self.assertEqual(error_points(rendered),rendered)
            self.assertTrue(all(re.match(r'（\d+）',line) for line in rendered.splitlines() if line.strip()))
        display=error_points('读数误差。\n$$u=\\sqrt{a^2+b^2}$$\n温度误差。')
        self.assertIn('\n\n$$u=\\sqrt{a^2+b^2}$$\n\n',display)
        self.assertEqual(error_points(display),display)
        self.assertIn('$u=1.23$',error_points('公式 $u=1.23$。其他来源。'))
        self.assertIn('%%DATA:x:%.2f%%',error_points('结果 %%DATA:x:%.2f%%。其他来源。'))
    def test_all_variants_are_numbered(self):
        for p in BASE.glob('*/variants.json'):
            for text in json.loads(p.read_text(encoding='utf-8')).get('误差分析',[]):
                with self.subTest(experiment=p.parent.name):
                    self.assertEqual(error_points(text),text)
                    self.assertGreaterEqual(len([v for v in text.splitlines() if v.strip()]),2)
    def test_all_reports_put_each_title_below_its_picture(self):
        for exp in all_experiments():
            with self.subTest(experiment=exp),tempfile.TemporaryDirectory() as tmp:
                photo=Path(tmp)/'record.png';Image.new('RGB',(120,80),'white').save(photo)
                with patch.dict(os.environ,{'LAB_DATA_PHOTO':str(photo)}),contextlib.redirect_stdout(io.StringIO()):out=Path(run_report(exp))
                try:
                    doc=Document(out);paras=doc.paragraphs
                    for i,p in enumerate(paras):
                        if p._p.find('.//'+qn('w:drawing')) is None:continue
                        self.assertFalse(p.text.strip(), 'image must have its own paragraph')
                        self.assertLess(i+1,len(paras));title=paras[i+1]
                        self.assertTrue(title.text.strip())
                        self.assertEqual(title.alignment,WD_ALIGN_PARAGRAPH.CENTER)
                        self.assertEqual(title.paragraph_format.first_line_indent.pt,0)
                        self.assertNotIn('注：',title.text)
                        self.assertNotIn('。',title.text)
                    for i,p in enumerate(paras):
                        if p.text.strip()=='误差分析':
                            points=[]
                            for line in paras[i+1:]:
                                if line.text.strip() in ('结论','实验结论') or '思考题' in line.text or re.match(r'^[一二三四五六七八九十]+、',line.text):break
                                if line.text.strip():points.append(line.text)
                            self.assertGreaterEqual(len(points),2)
                            self.assertTrue(all(re.match(r'（\d+）',v) for v in points),points)
                finally:shutil.rmtree(out.parent)
    def test_custom_figure_and_chart_insert_separate_notes(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);image=root/'chart.png';Image.new('RGB',(100,100),'white').save(image)
            out=root/'report.docx';writer=DocxReportWriter(str(out));writer.add_heading('实验结果分析')
            writer.add_image(str(image),caption='图1 曲线。灰点为实测数据。');writer.save();writer.close()
            doc=Document(out);self.assertEqual(doc.paragraphs[-2].text,'图1 曲线');self.assertEqual(doc.paragraphs[-1].text,'注：灰点为实测数据。')
            spec=importlib.util.spec_from_file_location('style_insert',ROOT/'src/main/insert-chart-into-docx.py');mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
            mod.insert_chart(str(out),str(image),'实验结果分析',caption='图2 新曲线\n注：虚线为拟合。')
            doc=Document(out)
            p=next(p for p in doc.paragraphs if p.text=='图2 新曲线');self.assertEqual(p.alignment,WD_ALIGN_PARAGRAPH.CENTER)
            self.assertIn('注：虚线为拟合。',[p.text for p in doc.paragraphs])
if __name__=='__main__':unittest.main()
