import json,sys,unittest,tempfile,os
from pathlib import Path
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'物理实验/实验脚本'))
from common.content_policy import inspect_text
from common.variants import render_custom_quiz
from docx import Document

class TieredIntegrityTests(unittest.TestCase):
 def test_shared_frontend_backend_cases(self):
  data=dict(r_mohm=94.0195875,r_0=150.25,c_x_bar=.0985972,focal_length=15.125)
  for case in json.loads((ROOT/'tests/tiered-integrity-cases.json').read_text(encoding='utf-8')):
   with self.subTest(text=case['text']):
    if case['status']=='block':
     with self.assertRaises(ValueError):inspect_text(case['text'],data)
    else:self.assertEqual(bool(inspect_text(case['text'],data)['warnings']),case['status']=='warn')
 def test_custom_answer_warning_renders_and_omitted_answer_stays_blank(self):
  class Writer:
   def __init__(self):self.doc=Document()
   def add_heading(self,*args,**kwargs):return self.doc.add_heading(*args,**kwargs)
   def add_paragraph_rich(self,text):return self.doc.add_paragraph(text)
   def add_paragraph(self,text):return self.doc.add_paragraph(text)
  quiz={'questions':['题目甲','题目乙'],'answers':{'2':['仪器精度为0.01 mV；结果%%DATA:r_mohm:%.2f%% MΩ。']}}
  with patch.dict(os.environ,{'LAB_JOB_INPUT':'','LAB_CUSTOM_QUIZ':json.dumps(quiz)}):
   writer=Writer();self.assertTrue(render_custom_quiz(writer,{'r_mohm':94.0195875}))
  with tempfile.TemporaryDirectory() as temp:
   output=Path(temp)/'report.docx';writer.doc.save(output)
   text='\n'.join(p.text for p in Document(output).paragraphs)
   self.assertIn('本次未生成回答',text);self.assertIn('0.01 mV',text);self.assertIn('94.02 MΩ',text)

if __name__=='__main__':unittest.main()
