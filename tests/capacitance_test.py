"""Independent analytic RC oracle, input failures, contamination gates and generation checks."""
import copy, importlib.util, json, math, os, shutil, sys, tempfile, unittest
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
EXP=ROOT/'物理实验/实验脚本/电容与高电阻的测量'
sys.path.insert(0,str(EXP));sys.path.insert(0,str(EXP.parent))
spec=importlib.util.spec_from_file_location('capacitance_generator',EXP/'generate.py');g=importlib.util.module_from_spec(spec);spec.loader.exec_module(g)
from report_integrity import validate_text
from common.variants import render_variant

def fixture():
    # Exact exponential with independently chosen parameters; no sample file as oracle.
    c, resistance=2.5,75.25
    t=[2*i for i in range(16)]
    return {'c_n':c,'r_0':123.75,'q_n':[2*(i+1) for i in range(8)],
            'q_x':[.32*2*(i+1)/c for i in range(8)],'t_values':t,
            'q_values':[12.3*math.exp(-v/(resistance*c)) for v in t]}

class CapacitanceTests(unittest.TestCase):
    def test_analytic_rc_and_capacitance(self):
        r=g._compute(fixture())
        self.assertAlmostEqual(r['c_x_bar'],.32,places=12)
        self.assertAlmostEqual(r['r_mohm'],75.25,places=10)
        self.assertAlmostEqual(r['r_fit_mohm'],75.25,places=10)
        self.assertAlmostEqual(r['k_slope'],-1/(75.25*2.5),places=14)
    def test_nominal_value_is_user_input(self):
        for nominal in [10,100,150,123.75]:
            d=fixture();d['r_0']=nominal;r=g._compute(d)
            self.assertAlmostEqual(r['delta_e'],abs(75.25/nominal-1)*100,places=10)
    def test_missing_nonfinite_and_nonpositive(self):
        for key in ['r_0','c_n']:
            for val in [None,0,-1,float('inf'),float('nan'),True,'100']:
                d=fixture();d[key]=val
                with self.subTest(key=key,val=val),self.assertRaises(ValueError):g._compute(d)
        for key in ['q_n','q_x','q_values']:
            d=fixture();d[key][0]=0
            with self.assertRaises(ValueError):g._compute(d)
    def test_arrays_time_and_slope(self):
        for key in ['q_n','q_x','t_values','q_values']:
            d=fixture();d[key]=d[key][:-1]
            with self.assertRaises(ValueError):g._compute(d)
        for val in [-1,0]:
            d=fixture();d['t_values'][1]=val
            with self.assertRaises(ValueError):g._compute(d)
        for q in [[1]*16,list(range(1,17))]:
            d=fixture();d['q_values']=q
            with self.assertRaisesRegex(ValueError,'负斜率'):g._compute(d)
    def test_blank_factory_data_and_no_defaults(self):
        for name in ['data.json','sample.json']:
            d=json.loads((EXP/name).read_text(encoding='utf-8'))
            self.assertTrue(all(v is None or all(x is None for x in v) for v in d.values()))
        schema=json.loads((EXP/'schema.json').read_text(encoding='utf-8'))
        for group in schema['groups']:
            for field in group['fields']:self.assertIsNone(field.get('default'))
    def test_all_variants_rebind(self):
        variants=json.loads((EXP/'variants.json').read_text(encoding='utf-8'))
        r=g._compute(fixture());d=fixture();d['r_0']=150;r2=g._compute(d)
        for texts in variants.values():
            for text in texts:
                validate_text(text,r)
                self.assertNotIn('%%DATA:',render_variant(text,r))
                if '%%DATA:delta_e:' in text:self.assertNotEqual(render_variant(text,r),render_variant(text,r2))
    def test_stale_overrides_and_unknown_placeholders_rejected(self):
        r=g._compute(fixture())
        for text in ['R=10.48 MΩ','放电时间为15.98 s','最长16秒','误差4.8%',r'$R=10\ \mathrm{MΩ}$','R_0=10','标称值为100','结果150','%%DATA:missing:%.2f%%','%%DATA:r_0:%s%%']:
            with self.subTest(text=text),self.assertRaises(ValueError):validate_text(text,r)
        validate_text(r'$R=1/(|K|C_N)$，结果%%DATA:r_mohm:%.2f%% MΩ',r)
    def test_rag_contains_no_sample_results(self):
        text=(EXP/'rag/原理.md').read_text(encoding='utf-8')
        for old in ['10.48','15.98','0.0954','16~17','16\\~17','<table>']:self.assertNotIn(old,text)
    def test_invalid_input_does_not_touch_old_report(self):
        with tempfile.TemporaryDirectory() as temp:
            p=Path(temp)/'old.docx';p.write_bytes(b'OLD_REPORT');d=fixture();d['r_0']=None
            with self.assertRaises(ValueError):g._generate_docx(d,str(p))
            self.assertEqual(p.read_bytes(),b'OLD_REPORT')
    def test_stale_polish_rejected_before_writing(self):
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as temp,patch.dict(os.environ,{'LAB_POLISH':json.dumps({'结论':'R=10.48 MΩ'}),'LAB_JOB_INPUT':''}):
            p=Path(temp)/'report.docx'
            with self.assertRaises(ValueError):g._generate_docx(fixture(),str(p))
            self.assertFalse(p.exists())

    def test_real_measurement_and_generated_word(self):
        from unittest.mock import patch
        from docx import Document
        d={'c_n':1,'r_0':100,
           'q_n':[2.02,4.04,6.09,8.11,10.18,12.24,14.30,16.37],
           'q_x':[.201,.399,.600,.800,1.002,1.204,1.406,1.608],
           't_values':[9.76,19.61,29.08,39.82,49.44,59.52,70,80,90.04,99.95,110.20,119.49,130.03,139.88,149.26,159.96],
           'q_values':[9.19,8.27,7.48,6.68,6.02,5.41,4.83,4.34,3.91,3.52,3.15,2.86,2.55,2.30,2.08,1.86]}
        result=g._compute(d)
        self.assertAlmostEqual(result['c_x_bar'],.09859719911282185,places=12)
        self.assertAlmostEqual(result['r_mohm'],94.01958751480782,places=10)
        self.assertAlmostEqual(result['r_fit_mohm'],93.95099562572071,places=10)
        with tempfile.TemporaryDirectory() as temp:
            for file in ['schema.json','variants.json']:shutil.copy2(EXP/file,Path(temp)/file)
            env={'LAB_JOB_INPUT':'','LAB_POLISH':'','LAB_CUSTOM_QUIZ':'','LAB_VARIANTS':json.dumps({key:0 for key in ['实验原理','实验方法','误差分析','结论']})}
            with patch.object(g,'SCRIPT_DIR',temp),patch.dict(os.environ,env):
                output=Path(temp)/'report.docx';g._generate_docx(d,str(output))
            from zipfile import ZipFile
            from lxml import etree
            with ZipFile(output) as z:
                xml=etree.fromstring(z.read('word/document.xml'))
                content=''.join(xml.xpath('//*[local-name()="t"]/text()'))
                for value in ['0.0986','94.02','5.98','93.95','6.05']:self.assertIn(value,content)
                self.assertNotIn('%%DATA:',content)
                for bad in ['10.48','4.8%','15.98 s','0.990']:self.assertNotIn(bad,content)
            doc=Document(output)
            self.assertEqual(doc.tables[0].rows[3].cells[1].text,'0.0995')

if __name__=='__main__':unittest.main()
