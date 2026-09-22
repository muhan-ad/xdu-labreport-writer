# -*- coding: utf-8 -*-
"""自定义画图（render_custom_plot）回归测试：内置图替换 / 无图实验追加图位。

覆盖两块：
1) 10 个有内置图的实验：注入 LAB_CUSTOM_PLOT 后内置图不再插入、图注被替换；
   不注入时内置图与图注照旧（回归）。
2) 17 个无内置图的实验：注入后在「数据处理」末尾出现自定义图；不注入时文档无图。

注入走 LAB_CUSTOM_PLOT 环境变量（等价于应用写进 LAB_JOB_INPUT 的同一字段），测试不写任何任务文件。
绘图运行器（src/main/plot-runner.py）的协议用例见 tests/plot_runner_test.py。
"""

import importlib.util
import json
import os
import re
import sys
import tempfile

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SCRIPTS = os.path.join(ROOT, "物理实验", "实验脚本")

# 有内置图的实验：slots=内置图位数，captions=内置图注前缀（用于断言被替换）
FIGURE_CASES = {
    "RLC电路的稳态特性研究实验": (2, []),
    "光的偏振特性测量": (1, []),
    "理想气体状态方程": (2, []),
    "电子偏转特性的测量": (3, ["图1", "图2", "图3"]),
    "电子元件伏安特性的测量": (3, []),
    "电容与高电阻的测量": (1, []),
    "直螺线管磁场分布的测量": (1, ["图1 螺线管"]),
    "重力加速度的测量": (1, []),
    "霍尔效应实验": (2, ["图1 B–Im", "图2 B–I"]),
    "静电场的模拟": (1, []),
}

_exp_cache = {}


def safe_join(base, *parts):
    """拼接路径并校验仍在 base 目录内（禁止越界，安全审查要求显式校验）。"""
    root = os.path.realpath(base)
    p = os.path.realpath(os.path.join(root, *parts))
    if os.path.commonpath([p, root]) != root:
        raise ValueError("路径越界：%s" % p)
    return p


def load_generator(exp):
    """把某实验的 generate.py 作为模块加载（不会执行 main()）。"""
    if exp in _exp_cache:
        return _exp_cache[exp]
    path = safe_join(SCRIPTS, exp, "generate.py")
    spec = importlib.util.spec_from_file_location("plot_gen_" + re.sub(r"\W+", "_", exp), path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    _exp_cache[exp] = mod
    return mod


def all_experiments():
    return sorted(n for n in os.listdir(SCRIPTS)
                  if os.path.isfile(safe_join(SCRIPTS, n, "generate.py")))


def make_png(path, color=(200, 60, 60)):
    """造一张纯色 PNG 当 AI 生成的图（尺寸小，插入 docx 很快）。"""
    from PIL import Image
    Image.new("RGB", (240, 160), color).save(path)


def run_report(exp, images=None):
    """按应用的方式生成一份报告到临时目录，返回 docx 路径。

    images: [{"path": ..., "caption": ...}] —— 模拟应用传入的自定义图；
            为 None 表示不启用自定义画图（走内置图）。
    """
    mod = load_generator(exp)
    with open(safe_join(SCRIPTS, exp, "data.json"), encoding="utf-8") as f:
        data = json.load(f)
    out = safe_join(tempfile.mkdtemp(prefix="custom_plot_"), "report.docx")
    orig = {k: os.environ.get(k) for k in ("LAB_CUSTOM_PLOT", "LAB_JOB_INPUT")}
    try:
        os.environ.pop("LAB_JOB_INPUT", None)      # 走 legacy 环境变量通路，测试不写任何任务文件
        if images is None:
            os.environ.pop("LAB_CUSTOM_PLOT", None)
        else:
            os.environ["LAB_CUSTOM_PLOT"] = json.dumps({"images": images}, ensure_ascii=False)
        mod._generate_docx(data, out)
    finally:
        for k, v in orig.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
    return out


def docx_paras(path):
    return [p.text.strip() for p in Document(path).paragraphs]


def image_count(path):
    return len(Document(path).inline_shapes)


def caption_is_own_paragraph(path, caption):
    """图注必须是独立段落（不能和图片挤在同一段里）。"""
    from docx.oxml.ns import qn
    for child in Document(path).element.body.iterchildren():
        if child.tag != qn('w:p'):
            continue
        text = ''.join(n.text or '' for n in child.iter(qn('w:t')))
        if caption in text:
            return not child.findall('.//' + qn('w:drawing'))
    return False


def main():
    problems = []
    tmp = tempfile.mkdtemp(prefix="plot_fixtures_")
    img1 = safe_join(tmp, "ai1.png")
    img2 = safe_join(tmp, "ai2.png")
    make_png(img1)
    make_png(img2, (60, 90, 200))
    custom = [{"path": img1, "caption": "自定义图注一"}, {"path": img2, "caption": "自定义图注二"}]

    exps = all_experiments()
    print("实验总数：%d（有内置图 %d 个）" % (len(exps), len(FIGURE_CASES)))

    # ---- 1) 有内置图的实验：替换 + 回归 ----
    for exp, (slots, builtin_captions) in sorted(FIGURE_CASES.items()):
        base = run_report(exp)
        if image_count(base) != slots:
            problems.append("%s：基线内置图 %d 张（应为 %d）—— 测试自身失效"
                            % (exp, image_count(base), slots))
            continue
        base_text = "\n".join(docx_paras(base))
        for cap in builtin_captions:
            if cap not in base_text:
                problems.append("%s：基线缺少内置图注「%s」" % (exp, cap))

        out = run_report(exp, custom)
        expect = min(len(custom), slots)
        if image_count(out) != expect:
            problems.append("%s：自定义画图后图 %d 张（应为 %d，内置图不得再插入）"
                            % (exp, image_count(out), expect))
        text = "\n".join(docx_paras(out))
        for cap in builtin_captions:
            if cap in text:
                problems.append("%s：自定义画图后仍出现内置图注「%s」" % (exp, cap))
        for item in custom[:expect]:
            if item["caption"] not in text:
                problems.append("%s：缺少自定义图注「%s」" % (exp, item["caption"]))
            elif not caption_is_own_paragraph(out, item["caption"]):
                problems.append("%s：图注「%s」与图片挤在同一段" % (exp, item["caption"]))
        print("   ✔ %s：基线 %d 张 → 自定义 %d 张" % (exp, slots, image_count(out)))

    # ---- 2) 无内置图的实验：追加图位 ----
    no_figure = [e for e in exps if e not in FIGURE_CASES]
    for exp in no_figure:
        if image_count(run_report(exp)) != 0:
            problems.append("%s：基线不应有图" % exp)
        out = run_report(exp, custom[:1])
        if image_count(out) != 1:
            problems.append("%s：自定义画图后应有 1 张图，实得 %d 张" % (exp, image_count(out)))
        elif custom[0]["caption"] not in "\n".join(docx_paras(out)):
            problems.append("%s：缺少自定义图注" % exp)
    print("   ✔ %d 个无内置图实验：不注入无图、注入后各 1 张" % len(no_figure))

    if problems:
        print("\n验证失败：")
        for p in problems:
            print("  - " + p)
        return 1
    print("\nPASS: 自定义画图替换内置图、无图实验追加图位")
    return 0


if __name__ == "__main__":
    sys.exit(main())
