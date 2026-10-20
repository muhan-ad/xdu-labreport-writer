"""Isolation, optical model, validation, measured plots and DOCX regression."""
import copy
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT/'物理实验/实验脚本'
NAME = '光的偏振特性测量（第二学期）'
EXP = BASE/NAME
sys.path.insert(0, str(BASE))
from common.polarization import compute, plot
from common.data_validation import validate
from common.variants import render_variant

class PolarizationTests(unittest.TestCase):
    def setUp(self):
        self.data = json.loads((EXP/'test-data.json').read_text(encoding='utf-8'))

    def test_fit_and_closure_have_independent_roles(self):
        result = compute(self.data)
        self.assertEqual((result['n_unique'], result['n_recorded']), (24,25))
        self.assertAlmostEqual(result['groups'][0]['A'], .61, places=6)
        self.assertAlmostEqual(result['groups'][0]['modulation'], .6/.61, places=6)
        self.assertAlmostEqual(result['groups'][1]['modulation'], .3/.61, places=5)
        self.assertLess(result['groups'][2]['modulation'], 1e-12)
        self.data['power_mw'][-1][0] += .123
        changed = compute(self.data)
        self.assertAlmostEqual(changed['groups'][0]['A'], result['groups'][0]['A'])
        self.assertAlmostEqual(changed['groups'][0]['closure'], .123)
        self.assertEqual(self.data['power_mw'][0][0], .01)  # No per-group minimum subtraction.

    def test_invalid_input_never_fills_defaults(self):
        for value in [None, -1, float('nan'), float('inf'), True, '1']:
            data = copy.deepcopy(self.data); data['power_mw'][7][1] = value
            with self.subTest(value=value), self.assertRaises(ValueError): compute(data)
        for rows in [[], self.data['power_mw'][:-1], self.data['power_mw']+[[1,2,3]], [[1,2]]*25]:
            with self.subTest(rows=len(rows)), self.assertRaises(ValueError): compute({'power_mw':rows})
        with self.assertRaises(ValueError): compute({'power_mw':[[0,0,0]]*25})
        schema = json.loads((EXP/'schema.json').read_text(encoding='utf-8'))
        self.assertFalse(validate(schema, json.loads((EXP/'data.json').read_text(encoding='utf-8')))['ok'])
        data=copy.deepcopy(self.data);data['wavelength_nm']=None;data['resolution_mw']=None
        self.assertEqual(compute(data)['groups'], compute(self.data)['groups'])

    def test_data_change_and_variants_are_bound(self):
        spec=importlib.util.spec_from_file_location('polarization_generate', EXP/'generate.py')
        module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        first=module._compute(self.data)
        self.data['power_mw'][6][0] += .1
        second=module._compute(self.data)
        self.assertNotEqual(first['m0'],second['m0'])
        self.assertNotEqual(first['summary'],second['summary'])
        for section,texts in json.loads((EXP/'variants.json').read_text(encoding='utf-8')).items():
            for text in texts:
                self.assertNotIn('%%DATA:',render_variant(text,second))

    def test_photo_suppresses_only_raw_tables_and_failure_keeps_them(self):
        from unittest.mock import patch
        from PIL import Image
        from docx import Document
        spec=importlib.util.spec_from_file_location('polar_photo_generate', EXP/'generate.py')
        module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary);photo=root/'record.png'
            Image.new('RGB',(100,100),'white').save(photo)
            for kind,source,expected_tables,expected_images in [
                ('valid',photo,1,2),('missing',root/'absent.png',3,1),('broken',root/'broken.png',3,1)]:
                if kind=='broken':source.write_bytes(b'invalid image')
                env={k:v for k,v in os.environ.items() if not k.startswith('LAB_')}
                env['LAB_DATA_PHOTO']=str(source)
                with self.subTest(kind=kind),patch.dict(os.environ,env,clear=True):
                    output=root/(kind+'.docx');module._generate_docx(self.data,str(output))
                    doc=Document(output)
                    self.assertEqual(len(doc.tables),expected_tables)
                    self.assertEqual(len(doc.inline_shapes),expected_images)
                    self.assertIn('A / mW',[c.text for c in doc.tables[-1].rows[0].cells])
                    self.assertIn('二、数据处理','\n'.join(p.text for p in doc.paragraphs))

    def test_report_and_preview_use_same_current_data(self):
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary)
            shutil.copytree(BASE/'common',root/'common',ignore=shutil.ignore_patterns('__pycache__'))
            exp=root/NAME;shutil.copytree(EXP,exp,ignore=shutil.ignore_patterns('__pycache__'))
            (exp/'data.json').write_text(json.dumps(self.data),encoding='utf-8')
            env={k:v for k,v in os.environ.items() if not k.startswith('LAB_')}
            generated=subprocess.run([sys.executable,'-B','-X','utf8',str(exp/'generate.py')],cwd=exp,env=env,capture_output=True,text=True,encoding='utf-8',timeout=60)
            self.assertEqual(generated.returncode,0,generated.stderr)
            docx=exp/(NAME+'实验报告.docx')
            from docx import Document
            doc=Document(docx)
            text='\n'.join(p.text for p in doc.paragraphs)
            self.assertIn(NAME,text);self.assertNotIn('马吕斯定律已被验证',text)
            self.assertNotIn('%%DATA:',text)
            self.assertEqual(len(doc.tables[0].rows)+len(doc.tables[1].rows)-2,25)
            self.assertEqual(len(doc.inline_shapes),1)
            with ZipFile(docx) as z:self.assertIn(b'm:oMath',z.read('word/document.xml'))
            source=(ROOT/'src/main/chart-preview.py').read_text(encoding='utf-8')
            output=root/'preview.png'
            preview=subprocess.run([sys.executable,'-B','-X','utf8','-', '--exp-path',str(exp),'--scripts-root',str(root),
                '--x-field','theta_deg','--y-field','power_mw','--chart-type','polar','--output',str(output)],
                input=source,cwd=exp,env=env,capture_output=True,text=True,encoding='utf-8',timeout=60)
            self.assertEqual(preview.returncode,0,preview.stderr)
            self.assertEqual(output.read_bytes(),(exp/'第二学期偏振光极坐标图.png').read_bytes())
            original=docx.read_bytes()
            self.data['power_mw'][0][0]=None
            (exp/'data.json').write_text(json.dumps(self.data),encoding='utf-8')
            bad=subprocess.run([sys.executable,'-B','-X','utf8',str(exp/'generate.py')],cwd=exp,env=env,capture_output=True,text=True,encoding='utf-8',timeout=60)
            self.assertNotEqual(bad.returncode,0)
            self.assertNotIn('报告已生成',bad.stdout)
            self.assertEqual(docx.read_bytes(),original)

if __name__=='__main__':unittest.main()
