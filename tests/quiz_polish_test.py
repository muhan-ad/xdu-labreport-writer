# -*- coding: utf-8 -*-
"""思考题章节回归：整段润色覆盖不得与写死的题目重复输出。

历史缺陷（2026-09-20 用户报告）：批量「导入润色结果并重新生成报告」把「思考题」写成整段
覆盖（str），generate.py 输出整段文本后又无条件输出写死题目 + 硬编码兜底答案 → 题目与
答案各出现两遍（12+ 个实验受影响）。修复在 common/variants.py：整段润色文本先解析成
「按问回答」，交给各脚本既有的「题目写死 + 按问取回答」路径输出。

本测试直接驱动各实验的 generate.py（与应用同一条链路，python-docx 读回，不需要 Word）。
润色覆盖通过 LAB_POLISH 环境变量注入（variants.py 里 LAB_JOB_INPUT 缺失时的等价通路），
因此测试不写任何文件：
  1) 全部含思考题的实验：整段覆盖模式下，题目行既不能重复、也不能丢失；
  2) 「声速的测量（空气）」四场景：整段覆盖 / 按问变体 / 无覆盖 / 解析失败（忽略 + 告警）。

用法：python -B -X utf8 tests/quiz_polish_test.py
"""
import contextlib
import importlib.util
import io
import json
import os
import re
import sys
import tempfile

from docx import Document

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(ROOT, "物理实验", "实验脚本")
if SCRIPTS not in sys.path:
    sys.path.insert(0, SCRIPTS)

QUIZ_SECTION_RE = re.compile(r"(思考题|问题讨论)")
# 应用自己的 docx 写入器不用 python-docx 的内置标题样式（段落样式全是 Normal），
# 因此按文本特征切小节：一级标题形如「四、课后思考题」。
SECTION_HEAD_RE = re.compile(r"^[一二三四五六七八九十]+、\s*\S")
QUIZ_HEAD_RE = re.compile(r"^[一二三四五六七八九十]+、\s*.*(思考题|问题讨论)")
Q_ITEM_RE = re.compile(r"^\s*(\d+)\.\s*\S")

_exp_cache = {}


def safe_join(base, *parts):
    """拼接路径并校验仍在 base 目录内（禁止越界，安全审查要求显式校验）。"""
    root = os.path.realpath(base)
    p = os.path.realpath(os.path.join(root, *parts))
    if os.path.commonpath([p, root]) != root:
        raise ValueError("路径越界：%s" % p)
    return p


def load_generator(exp):
    """把某实验的 generate.py 作为模块加载（不会执行 main()）。"""
    if exp in _exp_cache:
        return _exp_cache[exp]
    path = safe_join(SCRIPTS, exp, "generate.py")
    spec = importlib.util.spec_from_file_location("quiz_gen_" + re.sub(r"\W+", "_", exp), path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    _exp_cache[exp] = mod
    return mod


def experiments_with_quiz():
    """返回所有「消费 variants['思考题']」的实验名。"""
    out = []
    for name in sorted(os.listdir(SCRIPTS)):
        p = os.path.join(SCRIPTS, name, "generate.py")
        if not os.path.isfile(p):
            continue
        with open(p, encoding="utf-8", errors="replace") as f:
            if 'variants.get("思考题")' in f.read():
                out.append(name)
    return out


def run_report(exp, polish=None, variants_patch=None, custom_quiz=None):
    """按应用的方式生成一份报告到临时目录，返回 docx 路径。

    polish: {章节名: 整段文本}（经 LAB_POLISH 注入）；
    variants_patch: 覆盖 load_variants() 的返回值（模拟磁盘上的按问变体字典）；
    custom_quiz: {questions: [...], answers: {...}} —— 模拟应用传入的自定义思考题
                 （走 LAB_CUSTOM_QUIZ 环境变量，等价于应用写进 LAB_JOB_INPUT 的同一字段）。
    """
    from common import variants as variants_mod

    mod = load_generator(exp)
    data_path = safe_join(SCRIPTS, exp, "data.json")
    with open(data_path, encoding="utf-8") as f:
        data = json.load(f)
    out = os.path.join(tempfile.mkdtemp(prefix="quiz_polish_"), "report.docx")
    orig_load = variants_mod.load_variants
    orig_env = {k: os.environ.get(k) for k in ("LAB_POLISH", "LAB_JOB_INPUT", "LAB_CUSTOM_QUIZ")}
    try:
        if variants_patch is not None:
            variants_mod.load_variants = lambda script_dir: variants_patch
        os.environ.pop("LAB_JOB_INPUT", None)      # 走 legacy 环境变量通路，测试不写任何文件
        if custom_quiz is not None:
            os.environ["LAB_CUSTOM_QUIZ"] = json.dumps(custom_quiz, ensure_ascii=False)
        else:
            os.environ.pop("LAB_CUSTOM_QUIZ", None)
        if polish is None:
            os.environ.pop("LAB_POLISH", None)
        else:
            os.environ["LAB_POLISH"] = json.dumps(polish, ensure_ascii=False)
        mod._generate_docx(data, out)
    finally:
        variants_mod.load_variants = orig_load
        for k, v in orig_env.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
    return out


def quiz_block(path):
    """读回 docx，返回思考题小节的段落文本列表（按一级标题文本切段）。"""
    paras = [p.text.strip() for p in Document(path).paragraphs]
    start = None
    for i, t in enumerate(paras):
        if start is None:
            if QUIZ_HEAD_RE.match(t):
                start = i + 1
            continue
        if SECTION_HEAD_RE.match(t):     # 下一个一级标题 → 小节结束
            return paras[start:i]
    return paras[start:] if start is not None else []


def question_lines(path):
    """思考题小节里的题目行（形如「1. …」）。"""
    return [t.strip() for t in quiz_block(path) if Q_ITEM_RE.match(t)]


def section_heads(path):
    """全文一级标题（形如「三、课后思考题」），用于检查整段覆盖后没有章节消失。"""
    return [t for t in (p.text.strip() for p in Document(path).paragraphs) if SECTION_HEAD_RE.match(t)]


def docx_text(path):
    return "\n".join(p.text for p in Document(path).paragraphs)


def polished_text(questions):
    """按应用「整段润色」的形态合成文本：首行标题 + 每问「题目 + 答：…」。"""
    lines = ["思考题"]
    for i, q in enumerate(questions, 1):
        lines.append(q)
        lines.append("答：整段润色回答%d。" % i)
    return "\n".join(lines)


def main():
    problems = []
    exps = experiments_with_quiz()
    print("含思考题变体的实验：%d 个" % len(exps))

    # 1) 全量不变量：整段覆盖模式下题目不重复、也不丢失；章节不消失
    for exp in exps:
        base = run_report(exp)
        qs = question_lines(base)
        heads = section_heads(base)
        if not qs:
            problems.append("%s：基线报告里没找到题目行（测试自身失效）" % exp)
            continue
        out = run_report(exp, polish={"思考题": polished_text(qs)})
        text = docx_text(out)
        for i, q in enumerate(qs, 1):
            n = text.count(q)
            if n != 1:
                problems.append("%s：题目「%s」出现 %d 次（应为 1）" % (exp, q[:28], n))
            ans = "整段润色回答%d。" % i
            if text.count(ans) != 1:
                problems.append("%s：润色回答 %d 出现 %d 次（应为 1）" % (exp, i, text.count(ans)))
        lost = [h for h in heads if h not in section_heads(out)]
        if lost:
            problems.append("%s：整段覆盖后章节丢失 %s（多半是缩进进了上一问的 else）" % (exp, lost))
    print("① 整段覆盖不重复/不丢失：%s" % ("PASS" if not problems else "FAIL"))

    # 2) 声速的测量（空气）四场景
    exp = "声速的测量（空气）"
    qs = question_lines(run_report(exp))
    cases = []

    text = docx_text(run_report(exp, polish={"思考题": polished_text(qs)}))
    cases.append(("整段覆盖：题目各 1 次", all(text.count(q) == 1 for q in qs)))
    cases.append(("整段覆盖：润色答案各 1 次", all(text.count("整段润色回答%d。" % i) == 1 for i in range(1, len(qs) + 1))))
    cases.append(("整段覆盖：不再出现硬编码兜底答案", "相同：都是基于声波在空气中传播的原理进行测量" not in text))

    text = docx_text(run_report(exp, variants_patch={"思考题": {"1": ["按问回答一"], "2": ["按问回答二"]}}))
    cases.append(("按问变体：题目各 1 次 + 回答各 1 次",
                  all(text.count(q) == 1 for q in qs) and text.count("按问回答一") == 1 and text.count("按问回答二") == 1))

    text = docx_text(run_report(exp))
    cases.append(("无覆盖：题目各 1 次 + 兜底答案在",
                  all(text.count(q) == 1 for q in qs) and "相同：都是基于声波在空气中传播的原理进行测量" in text))

    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        text = docx_text(run_report(exp, polish={"思考题": "这段文字里没有任何题号，无法解析。"}))
    cases.append(("解析失败：忽略覆盖、题目各 1 次、不出现原文",
                  all(text.count(q) == 1 for q in qs)
                  and "这段文字里没有任何题号，无法解析。" not in text
                  and "quiz-polish" in err.getvalue()))

    # 3) 自定义思考题（应用侧传入）：内置题目被完全替换，答案用应用给的，章节结构不变
    base_heads = section_heads(run_report(exp))
    custom_qs = ["自定义题目一：为什么？", "自定义题目二：如何？", "自定义题目三：可否？"]
    text = docx_text(run_report(exp, custom_quiz={
        "questions": custom_qs,
        "answers": {"1": ["AI 回答一。"], "2": ["AI 回答二。"], "3": ["AI 回答三。"]},
    }))
    cases.append(("自定义题目：内置题目全部消失", all(q not in text for q in qs)))
    cases.append(("自定义题目：题面各 1 次且带题号",
                  all(text.count("%d. %s" % (i, q)) == 1 for i, q in enumerate(custom_qs, 1))))
    cases.append(("自定义题目：答案各 1 次",
                  all(text.count("AI 回答%s。" % c) == 1 for c in ("一", "二", "三"))))
    cases.append(("自定义题目：内置兜底答案不出现", "相同：都是基于声波在空气中传播的原理进行测量" not in text))
    cases.append(("自定义题目：章节结构不变", section_heads(run_report(exp, custom_quiz={
        "questions": custom_qs, "answers": {"1": ["AI 回答一。"]}})) == base_heads))
    cases.append(("自定义题目：题面仍能被题目行正则识别", len(question_lines(run_report(exp, custom_quiz={
        "questions": custom_qs, "answers": {}}))) == len(custom_qs)))
    text = docx_text(run_report(exp, custom_quiz={"questions": custom_qs, "answers": {}}))
    cases.append(("自定义题目：答案缺失时给占位、不报错",
                  text.count("答：（本次未生成回答）") == len(custom_qs)
                  and all(q in text for q in custom_qs)))

    for name, ok in cases:
        print("   %s %s" % ("✔" if ok else "✘", name))
        if not ok:
            problems.append("声速的测量（空气）：" + name)

    if problems:
        print("\n验证失败：")
        for p in problems:
            print("  - " + p)
        return 1
    print("\nPASS: 思考题整段覆盖不再重复输出，题目不丢失")
    return 0


if __name__ == "__main__":
    sys.exit(main())
